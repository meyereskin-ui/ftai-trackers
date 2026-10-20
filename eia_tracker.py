"""EIA-860M monitor for FTAI Power-sized gas turbines.

FTAI Power's unit is a CFM56 core repackaged as a ~25 MW gas turbine. EIA's monthly
generator inventory (EIA-860M) lists every grid-reported US generator, planned or
operating, with size and technology. This tracker keeps the natural-gas simple-cycle
turbines (prime mover GT) of 20-32 MW and alerts when new ones appear, flagging
clusters at one plant and data-center-sounding owners.

Limit: behind-the-meter plants that never report to EIA (xAI's Southaven turbines
are absent) won't show up. Treat a hit as a lead to check against permits and news.

EIA posts a new file about once a month; on other days the run is a no-op.
Usage: python eia_tracker.py run [--xlsx file.xlsx]
"""
from __future__ import annotations

import argparse
import io
import json
import re
from datetime import date

import openpyxl
import requests

from common import BROWSER_UA, connect, now_utc, record_alerts

PAGE = "https://www.eia.gov/electricity/data/eia860m/"
MW_MIN, MW_MAX = 20.0, 32.0
CLUSTER_MIN = 3
DC_HINTS = ("DATA", "COMPUTE", "DIGITAL", "CLOUD", "AI ", "XAI", "HYPERSCALE", "TECHNOLOG",
            "CRUSOE", "CORE SCIENTIFIC", "SILICON", "COLOSSUS", "STARGATE", "MINING")
SHEETS = {"Operating": "operating", "Planned": "planned"}


def candidate_files(html: str) -> list[str]:
    links = re.findall(r'href="([^"]*_generator(\d{4})\.xlsx)"', html)
    months = ["january", "february", "march", "april", "may", "june", "july", "august",
              "september", "october", "november", "december"]

    def key(link):
        m = re.search(r"/(\w+)_generator(\d{4})\.xlsx", link)
        return (int(m.group(2)), months.index(m.group(1))) if m and m.group(1) in months else (0, 0)
    return sorted({l for l, _ in links}, key=key, reverse=True)


def fetch_latest(s: requests.Session) -> tuple[str, bytes]:
    """Newest file that actually has data (EIA posts empty placeholders ahead of time)."""
    for link in candidate_files(s.get(PAGE, timeout=60).text)[:6]:
        r = s.get("https://www.eia.gov" + link if link.startswith("/") else PAGE + link, timeout=300)
        if r.ok and r.content[:2] == b"PK":
            return link.rsplit("/", 1)[1], r.content
    raise RuntimeError("no EIA-860M file with data found")


def parse(blob: bytes) -> tuple[str, dict[str, dict]]:
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
    as_of, out = "", {}
    for sheet, status in SHEETS.items():
        rows = wb[sheet].iter_rows(values_only=True)
        title = next(rows)
        as_of = as_of or next((str(x) for x in title if x), "")
        header = None
        for r in rows:
            if header is None:
                if r and "Plant ID" in r:
                    header = {k: i for i, k in enumerate(r) if k}
                continue
            if not r or r[header["Plant ID"]] in (None, ""):
                continue
            g = lambda k: r[header[k]] if k in header else None
            try:
                mw = float(g("Nameplate Capacity (MW)") or 0)
            except ValueError:
                continue
            if g("Prime Mover Code") != "GT" or g("Energy Source Code") != "NG" or not MW_MIN <= mw <= MW_MAX:
                continue
            year = g("Planned Operation Year") or g("Operating Year")
            if status == "operating" and (not str(year).isdigit() or int(year) < date.today().year - 1):
                continue  # old peakers are not FTAI units
            key = f"{g('Plant ID')}|{g('Generator ID')}"
            out[key] = {"sheet": status, "entity": g("Entity Name"), "plant": g("Plant Name"),
                        "plant_id": g("Plant ID"), "gen": g("Generator ID"), "state": g("Plant State"),
                        "county": g("County"), "mw": mw, "year": year,
                        "month": g("Planned Operation Month") or g("Operating Month"),
                        "status": g("Status"), "lat": g("Latitude"), "lon": g("Longitude")}
    return as_of, out


def cmd_run(args) -> None:
    con = connect()
    con.executescript("""
        CREATE TABLE IF NOT EXISTS eia_units (key TEXT PRIMARY KEY, data TEXT, first_seen TEXT, as_of TEXT);
        CREATE TABLE IF NOT EXISTS eia_files (name TEXT PRIMARY KEY, processed TEXT);""")
    if args.xlsx:
        name, blob = args.xlsx.rsplit("/", 1)[-1], open(args.xlsx, "rb").read()
    else:
        s = requests.Session()
        s.headers.update({"User-Agent": BROWSER_UA})
        name, blob = fetch_latest(s)
    if con.execute("SELECT 1 FROM eia_files WHERE name=?", (name,)).fetchone():
        print(f"[eia] {name} already processed; nothing new this month")
        return
    as_of, units = parse(blob)
    baseline = con.execute("SELECT COUNT(*) FROM eia_files").fetchone()[0] == 0
    known = {r[0]: json.loads(r[1]) for r in con.execute("SELECT key, data FROM eia_units")}
    new = {k: u for k, u in units.items() if k not in known}
    moved = {k: u for k, u in units.items() if k in known and known[k]["sheet"] != u["sheet"]}
    con.executemany("INSERT OR REPLACE INTO eia_units VALUES (?,?,COALESCE((SELECT first_seen FROM eia_units WHERE key=?),?),?)",
                    [(k, json.dumps(u, default=str), k, now_utc(), as_of) for k, u in units.items()])
    con.execute("INSERT INTO eia_files VALUES (?,?)", (name, now_utc()))
    con.commit()
    print(f"[eia] {as_of}: {len(units)} gas turbines of {MW_MIN:g}-{MW_MAX:g} MW "
          f"({len(new)} new, {len(moved)} changed status)")
    if baseline:
        print("[eia] first file stored as the baseline; alerts start with the next monthly file")
        return
    plants: dict = {}
    for u in units.values():
        plants.setdefault(u["plant_id"], []).append(u)
    alerts = []
    for kind, group in (("NEW TURBINES", new), ("NOW OPERATING", moved)):
        by_plant: dict = {}
        for u in group.values():
            by_plant.setdefault(u["plant_id"], []).append(u)
        for pid, us in by_plant.items():
            u0 = us[0]
            size = len(plants.get(pid, us))
            dc = any(h in f"{u0['entity']} {u0['plant']}".upper() for h in DC_HINTS)
            hit = dc or size >= CLUSTER_MIN
            mw = sum(u["mw"] for u in us)
            maplink = f" · map https://maps.google.com/?q={u0['lat']},{u0['lon']}" if u0["lat"] else ""
            alerts.append({"kind": kind, "key": f"{pid}", "text":
                f"{'🔴' if hit else '⚪'} {u0['plant']} ({u0['entity']}) · {u0['county']} County, {u0['state']} · "
                f"{len(us)} unit(s), {mw:.0f} MW, {size} Power-sized units at the plant · "
                f"{u0['status']} · online {u0['month'] or '?'}/{u0['year'] or '?'}"
                f"{' · data-center-sounding owner' if dc else ''}"
                f"{maplink}"})
    record_alerts(con, "eia", alerts)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--xlsx", help="local EIA-860M file instead of downloading")
    r.set_defaults(func=cmd_run)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
